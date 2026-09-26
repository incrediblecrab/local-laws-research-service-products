"""The Census adapter against real rows of the 2022 file: normalization, the check against CG2200ORG02, and each way a changed source is refused."""

import pytest

from conftest import ORG02, UNITS, edited_units, sha256, units_rows
from local_laws import census
from local_laws.census import SourceChanged
from local_laws.schema import GOVERNMENTS


def test_the_sample_parses_with_each_normalization_counted():
    rows, notes = units_rows()
    assert len(rows) == 50
    assert notes == {"census_gid_missing": 1, "census_id_pid6_stored_as_number": 1, "dependent_school_systems_skipped": 1, "mailing_state_differs": 1, "strings_trimmed": 4}
    by_id = {row["census_id"]: row for row in rows}
    assert [row["census_id"] for row in rows] == sorted(by_id), "rows are sorted by census_id"
    assert by_id["136739"]["name"] == "TOWN OF GREENVILLE", "a PID6 stored as a number is read as its 6 digits"
    assert by_id["251235"]["name"] == "TOWN OF MULE BARN" and by_id["201532"]["name"] == "PLATEAU VALLEY SCH DIST 50", "surrounding spaces are trimmed"
    assert by_id["162028"]["web_address"] is None, "a web address of one space is null"
    assert by_id["100413"]["state"] == "AR", "state comes from FIPS_STATE, not the mailing address in Texas"
    assert by_id["248752"]["is_active"] is False and by_id["100001"]["is_active"] is True
    assert by_id["251419"]["census_gid"] is None
    assert by_id["102163"]["government_type"] == "municipal" and by_id["161769"]["government_type"] == "county", "COUNTY OF MACON-BIBB is a consolidated municipal government"
    assert by_id["122663"]["web_address"] == "https://www.mobilehousing.org/"
    assert "225903" not in by_id, "dependent school systems are not governments"
    assert by_id["177885"]["enrollment"] and by_id["177885"]["population"] is None and by_id["100001"]["enrollment"] is None
    assert all(list(row) == GOVERNMENTS.names for row in rows)


def test_the_sample_agrees_with_its_cg2200org02_and_a_planted_wrong_count_is_named():
    rows, _ = units_rows()
    org02 = census.parse_org02(census.zip_member(ORG02, census.ORG02_MEMBER))
    assert len(org02) == 312 and census.compare(rows, org02) == []
    org02[("50", "township")] += 1
    assert census.compare(rows, org02) == ["VT township: 1 rows, CG2200ORG02 says 2"]


def test_a_changed_file_is_refused_by_its_pin():
    assert census.check_sha256(UNITS, sha256(UNITS), census.GOVT_UNITS_URL) == sha256(UNITS)
    with pytest.raises(SourceChanged, match="not the pinned"):
        census.check_sha256(UNITS + b"x", sha256(UNITS), census.GOVT_UNITS_URL)


def planted(edit):
    return census.parse_units(census.zip_member(edited_units(edit), census.GOVT_UNITS_MEMBER))


def set_cell(sheet, census_id, column, value):
    def edit(book):
        rows = list(book[sheet].iter_rows())
        header = [cell.value for cell in rows[0]]
        for row in rows[1:]:
            if str(row[0].value).strip() == census_id:
                row[header.index(column)].value = value
                return
        raise AssertionError(f"{census_id} not in {sheet}")
    return edit


@pytest.mark.parametrize("edit, message", [
    (set_cell("General Purpose", "100001", "UNIT_TYPE", "4 - BOROUGH"), "unknown UNIT_TYPE"),
    (set_cell("General Purpose", "100019", "CENSUS_ID_PID6", "100001"), "duplicate census ids"),
    (set_cell("General Purpose", "100001", "FIPS_STATE", "72"), "not a state or DC"),
    (set_cell("Special District", "100117", "IS_ACTIVE", "maybe"), "not Y or N"),
    (set_cell("School District", "177885", "ENROLLMENT", "many"), "not an integer"),
    (set_cell("General Purpose", "100001", "FIPS_PLACE", "123"), "not 5 digits"),
    (lambda book: book.create_sheet("Tribal"), "unexpected sheets"),
    (lambda book: book["Special District"].delete_cols(4), "lacks"),
])
def test_each_planted_change_to_the_source_stops_the_parse(edit, message):
    with pytest.raises(SourceChanged, match=message):
        planted(edit)


def test_cg2200org02_must_cover_every_state():
    import io
    import zipfile

    import openpyxl

    book = openpyxl.load_workbook(io.BytesIO(census.zip_member(ORG02, census.ORG02_MEMBER)))
    sheet = book.worksheets[0]
    for row in reversed(list(sheet.iter_rows(min_row=2))):
        if row[9].value == "56":
            sheet.delete_rows(row[0].row)
    buffer, out = io.BytesIO(), io.BytesIO()
    book.save(buffer)
    with zipfile.ZipFile(out, "w") as archive:
        archive.writestr("COG2022_CG2200ORG02_Data.xlsx", buffer.getvalue())
    with pytest.raises(SourceChanged, match="covers 50 states, not the 50 and DC"):
        census.parse_org02(census.zip_member(out.getvalue(), census.ORG02_MEMBER))


def test_zip_member_needs_exactly_one_match():
    with pytest.raises(SourceChanged, match="expected one member"):
        census.zip_member(UNITS, r"nothing\.xlsx")
