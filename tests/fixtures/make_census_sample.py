"""Writes the Census fixtures from the real files: python tests/fixtures/make_census_sample.py path/to/govt_units_2022.ZIP

govt_units_sample.zip holds the chosen rows of the real Govt_Units_2022_Final.xlsx, copied value for value, with the contact title and the street, city and ZIP of the mailing address blanked: the dataset leaves them out.
org02_sample.zip holds a CG2200ORG02 computed from those rows, so the build's check against it passes; a test plants a wrong count to see it fail.
"""

import io
import sys
import zipfile
from pathlib import Path

import openpyxl

sys.path.insert(0, str(Path(__file__).resolve().parents[2]))

from local_laws import census  # noqa: E402

HERE = Path(__file__).resolve().parent
PRIVATE = {"TITLE", "ADDRESS1", "ADDRESS2", "CITY", "ZIP", "ZIP4"}
# Same-name governments the matching must tell apart, the governments the aliases name, and rows each normalization is for.
CHOSEN = {
    "General Purpose": [
        "100001", "100019", "114404", "135964", "109507", "170831", "125103", "162180", "109173", "170293", "170302", "186570", "127039", "127088",
        "175544", "187490", "240470", "177560", "201942", "166904", "123750", "185811", "166435", "102163", "161769", "194463", "166017",
        "207957", "161958", "165831", "166076", "175728", "127573", "170895", "164335", "126082",
        "136739", "251235", "162028", "100413", "248752",
    ],
    "Special District": ["100117", "158686", "251419", "122664", "122663"],
    "School District": ["177885", "117827", "180539", "201532"],
    "DEP School Dist": ["225903"],
}
ORG02_COLUMNS = ("GEO_ID", "GEO_TTL", "YEAR", "AMOUNT", "AGG_DESC", "AGG_DESC_TTL", "GOVTYPE", "GOVTYPE_TTL", "GEOTYPE", "ST", "SVY_COMP", "SVY_COMP_TTL")
ORG02_TITLES = {"GO0002": "Total Local Government Units", "GO0004": "Total Local Government Units - General Purpose Governments", "GO0005": "Total Local Government Units - County Governments",
                "GO0007": "Total Local Government Units - Subcounty Governments - Municipal Governments", "GO0008": "Total Local Government Units - Subcounty Governments - Township Governments",
                "GO0009": "Total Local Government Units - Special Purpose Governments - Special District Governments",
                "GO0010": "Total Local Government Units - Special Purpose Governments - Independent School District Governments"}


def zipped(member, book):
    buffer = io.BytesIO()
    book.save(buffer)
    out = io.BytesIO()
    with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as archive:
        archive.writestr(member, buffer.getvalue())
    return out.getvalue()


def sample_units(source):
    real = openpyxl.load_workbook(io.BytesIO(census.zip_member(source, census.GOVT_UNITS_MEMBER)), read_only=True, data_only=True)
    book = openpyxl.Workbook()
    book.remove(book.active)
    for sheet in real.sheetnames:
        rows = real[sheet].iter_rows(values_only=True)
        header = next(rows)
        wanted, found = set(CHOSEN[sheet]), {}
        for values in rows:
            uid = str(values[0]).strip() if values[0] is not None else None
            if uid in wanted:
                if uid in found:
                    raise SystemExit(f"{uid} appears twice in {sheet}")
                found[uid] = [None if column in PRIVATE else value for column, value in zip(header, values)]
        if set(found) != wanted:
            raise SystemExit(f"{sheet}: not found {sorted(wanted - set(found))}")
        target = book.create_sheet(sheet)
        target.append(header)
        for uid in CHOSEN[sheet]:
            target.append(found[uid])
    real.close()
    return zipped("Govt_Units_2022_Final.xlsx", book)


def sample_org02(units):
    rows, _ = census.parse_units(census.zip_member(units, census.GOVT_UNITS_MEMBER))
    counts = census.tally(rows)
    book = openpyxl.Workbook()
    sheet = book.active
    sheet.title = "CG2200ORG02"
    sheet.append(ORG02_COLUMNS)
    for state in [census.NATION] + sorted(census.STATES):
        for code, title in ORG02_TITLES.items():
            if code == "GO0004":
                amount = sum(counts.get((state, kind), 0) for kind in ("county", "municipal", "township"))
            else:
                amount = counts.get((state, census.ORG02_CODES[code]), 0)
            name = "United States" if state == census.NATION else census.STATES[state]
            sheet.append([f"0{'1' if state == census.NATION else '4'}00000US{'' if state == census.NATION else state}", name, "2022", amount or "X", code, title, "001", "State And Local",
                          "01" if state == census.NATION else "02", state, "07", "Government Organization"])
    return zipped("COG2022_CG2200ORG02_Data.xlsx", book)


def main(path):
    source = Path(path).read_bytes()
    census.check_sha256(source, census.GOVT_UNITS_SHA256, path)
    units = sample_units(source)
    (HERE / "govt_units_sample.zip").write_bytes(units)
    (HERE / "org02_sample.zip").write_bytes(sample_org02(units))
    print(f"wrote {HERE / 'govt_units_sample.zip'} and {HERE / 'org02_sample.zip'}")


if __name__ == "__main__":
    main(sys.argv[1])
