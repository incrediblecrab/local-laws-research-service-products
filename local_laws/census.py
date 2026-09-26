"""The 2022 Census of Governments: every local government in the 50 states and DC, and the Census's own count of them by type and state."""

import hashlib
import io
import re
import zipfile
from collections import Counter

import openpyxl

GOVT_UNITS_URL = "https://www2.census.gov/programs-surveys/gus/datasets/2022/govt_units_2022.ZIP"
ORG02_URL = "https://www2.census.gov/programs-surveys/gus/tables/2022/cog2022_cg2200org02.zip"
# SHA-256 of each file as downloaded on September 25, 2026. The Census published both in 2023; different bytes mean the source changed, which needs a person to look before anything is published from it.
GOVT_UNITS_SHA256 = "3ef3d93c91697b00d4d53384176dc584db9dacbd0f0a23e45657f0a3a1058c28"
ORG02_SHA256 = "7c34fe70ed9b05fcb2851cf809823b143301e34a1444007b044ac6f152481d48"
GOVT_UNITS_MEMBER = r"Govt_Units_2022_Final\.xlsx"
ORG02_MEMBER = r"COG2022_CG2200ORG02_Data\.xlsx"

TYPES = ("county", "municipal", "township", "special_district", "school_district")
GENERAL_PURPOSE = {"1 - COUNTY": "county", "2 - MUNICIPAL": "municipal", "3 - TOWNSHIP": "township"}
SHEETS = ("General Purpose", "Special District", "School District")
# Dependent public school systems belong to another government, so the Census does not count them as governments, and neither does this dataset.
SKIPPED_SHEETS = ("DEP School Dist",)
COMMON = {"CENSUS_ID_PID6", "CENSUS_ID_GIDID", "UNIT_NAME", "STATE", "WEB_ADDRESS", "FIPS_STATE", "FIPS_COUNTY", "COUNTY_AREA_NAME", "IS_ACTIVE"}
REQUIRED = {
    "General Purpose": COMMON | {"UNIT_TYPE", "POPULATION", "POPULATION_YEAR", "FIPS_PLACE"},
    "Special District": COMMON | {"FUNCTION_NAME"},
    "School District": COMMON | {"ENROLLMENT", "ENROLLMENT_YEAR", "SCHOOL_LEVEL_DESCRIPTION"},
}
# CG2200ORG02's aggregate codes, from the table's "Aggregate Descriptions" file.
ORG02_CODES = {"GO0002": "total", "GO0005": "county", "GO0007": "municipal", "GO0008": "township", "GO0009": "special_district", "GO0010": "school_district"}
NATION = "00"
STATES = {
    "01": "AL", "02": "AK", "04": "AZ", "05": "AR", "06": "CA", "08": "CO", "09": "CT", "10": "DE", "11": "DC", "12": "FL",
    "13": "GA", "15": "HI", "16": "ID", "17": "IL", "18": "IN", "19": "IA", "20": "KS", "21": "KY", "22": "LA", "23": "ME",
    "24": "MD", "25": "MA", "26": "MI", "27": "MN", "28": "MS", "29": "MO", "30": "MT", "31": "NE", "32": "NV", "33": "NH",
    "34": "NJ", "35": "NM", "36": "NY", "37": "NC", "38": "ND", "39": "OH", "40": "OK", "41": "OR", "42": "PA", "44": "RI",
    "45": "SC", "46": "SD", "47": "TN", "48": "TX", "49": "UT", "50": "VT", "51": "VA", "53": "WA", "54": "WV", "55": "WI",
    "56": "WY",
}


class SourceChanged(RuntimeError):
    """A source file is not the one this code was written and checked against."""


def check_sha256(data, expected, url):
    actual = hashlib.sha256(data).hexdigest()
    if actual != expected:
        raise SourceChanged(f"{url} has SHA-256 {actual}, not the pinned {expected}: review the new file, then update the pin")
    return actual


def zip_member(data, pattern):
    with zipfile.ZipFile(io.BytesIO(data)) as archive:
        names = [name for name in archive.namelist() if re.fullmatch(pattern, name.rsplit("/", 1)[-1])]
        if len(names) != 1:
            raise SourceChanged(f"expected one member matching {pattern}, found {names}")
        return archive.read(names[0])


def _text(value, notes):
    if value is None:
        return None
    value = str(value)
    stripped = value.strip()
    if stripped != value:
        notes["strings_trimmed"] += 1
    return stripped or None


def _code(value, width, field, notes, required=True):
    if value is None:
        if required:
            raise SourceChanged(f"{field} is missing")
        return None
    if isinstance(value, int):
        notes[f"{field.lower()}_stored_as_number"] += 1
        value = str(value).zfill(width)
    value = value.strip()
    if not re.fullmatch(rf"\d{{{width}}}", value):
        raise SourceChanged(f"{field} {value!r} is not {width} digits")
    return value


def _int(value, field):
    if value is None:
        return None
    if isinstance(value, bool) or not isinstance(value, int):
        raise SourceChanged(f"{field} {value!r} is not an integer")
    return value


def normalize(record, sheet, notes):
    """One governments row from one row of the Census file."""
    state_fips = _code(record["FIPS_STATE"], 2, "FIPS_STATE", notes)
    if state_fips not in STATES:
        raise SourceChanged(f"FIPS_STATE {state_fips} is not a state or DC")
    state = STATES[state_fips]
    mailing_state = _text(record["STATE"], notes)
    if mailing_state != state:
        notes["mailing_state_differs"] += 1
    if sheet == "General Purpose":
        kind = GENERAL_PURPOSE.get(record["UNIT_TYPE"])
        if kind is None:
            raise SourceChanged(f"unknown UNIT_TYPE {record['UNIT_TYPE']!r}")
    else:
        kind = "special_district" if sheet == "Special District" else "school_district"
    active = record["IS_ACTIVE"]
    if active not in ("Y", "N"):
        raise SourceChanged(f"IS_ACTIVE {active!r} is not Y or N")
    name = _text(record["UNIT_NAME"], notes)
    if not name:
        raise SourceChanged(f"government {record['CENSUS_ID_PID6']} has no name")
    return {
        "census_id": _code(record["CENSUS_ID_PID6"], 6, "CENSUS_ID_PID6", notes),
        "census_gid": _code(record["CENSUS_ID_GIDID"], 14, "CENSUS_ID_GIDID", notes, required=False),
        "name": name,
        "government_type": kind,
        "state": state,
        "state_fips": state_fips,
        "county_fips": _code(record["FIPS_COUNTY"], 3, "FIPS_COUNTY", notes, required=False),
        "county_name": _text(record["COUNTY_AREA_NAME"], notes),
        "fips_place": _code(record["FIPS_PLACE"], 5, "FIPS_PLACE", notes) if sheet == "General Purpose" else None,
        "special_district_function": _text(record.get("FUNCTION_NAME"), notes),
        "school_level": _text(record.get("SCHOOL_LEVEL_DESCRIPTION"), notes),
        "population": _int(record.get("POPULATION"), "POPULATION"),
        "population_year": _int(record.get("POPULATION_YEAR"), "POPULATION_YEAR"),
        "enrollment": _int(record.get("ENROLLMENT"), "ENROLLMENT"),
        "enrollment_year": _int(record.get("ENROLLMENT_YEAR"), "ENROLLMENT_YEAR"),
        "web_address": _text(record["WEB_ADDRESS"], notes),
        "is_active": active == "Y",
    }


def parse_units(xlsx):
    """The governments table's rows, sorted by census_id, and a count of each normalization applied."""
    book = openpyxl.load_workbook(io.BytesIO(xlsx), read_only=True, data_only=True)
    if sorted(book.sheetnames) != sorted(SHEETS + SKIPPED_SHEETS):
        raise SourceChanged(f"unexpected sheets {book.sheetnames}")
    rows, notes = [], Counter()
    for sheet in SHEETS:
        records = book[sheet].iter_rows(values_only=True)
        header = next(records)
        missing = REQUIRED[sheet] - set(header)
        if missing:
            raise SourceChanged(f"sheet {sheet!r} lacks {sorted(missing)}")
        for values in records:
            if all(value is None for value in values):
                notes["blank_rows_skipped"] += 1
                continue
            rows.append(normalize(dict(zip(header, values)), sheet, notes))
    notes["dependent_school_systems_skipped"] = sum(1 for values in book["DEP School Dist"].iter_rows(min_row=2, values_only=True) if any(value is not None for value in values))
    book.close()
    duplicates = sorted(uid for uid, n in Counter(row["census_id"] for row in rows).items() if n > 1)
    if duplicates:
        raise SourceChanged(f"duplicate census ids {duplicates[:10]}")
    rows.sort(key=lambda row: row["census_id"])
    notes["census_gid_missing"] = sum(1 for row in rows if row["census_gid"] is None)
    return rows, dict(sorted(notes.items()))


def parse_org02(xlsx):
    """Table CG2200ORG02, Local Governments by Type and State: 2022, as {(state_fips, type): count}. The nation is state "00"; "X" (not applicable) counts 0."""
    book = openpyxl.load_workbook(io.BytesIO(xlsx), read_only=True, data_only=True)
    records = book.worksheets[0].iter_rows(values_only=True)
    header = next(records)
    missing = {"ST", "AGG_DESC", "AMOUNT"} - set(header)
    if missing:
        raise SourceChanged(f"CG2200ORG02 lacks {sorted(missing)}")
    counts = {}
    for values in records:
        record = dict(zip(header, values))
        kind = ORG02_CODES.get(record["AGG_DESC"])
        if kind is None:
            continue
        amount = record["AMOUNT"]
        counts[(record["ST"], kind)] = 0 if amount == "X" else _int(amount, "AMOUNT")
    book.close()
    states = {state for state, _ in counts} - {NATION}
    if states != set(STATES):
        raise SourceChanged(f"CG2200ORG02 covers {len(states)} states, not the 50 and DC")
    return counts


def tally(rows):
    """What CG2200ORG02 counts, computed from governments rows."""
    counts = Counter()
    for row in rows:
        for state in (row["state_fips"], NATION):
            counts[(state, row["government_type"])] += 1
            counts[(state, "total")] += 1
    return counts


def compare(rows, org02):
    """Every cell of CG2200ORG02 that the rows disagree with, as text; empty when all agree."""
    counts = tally(rows)
    return [f"{STATES.get(state, 'US')} {kind}: {counts.get((state, kind), 0):,} rows, CG2200ORG02 says {expected:,}"
            for (state, kind), expected in sorted(org02.items()) if counts.get((state, kind), 0) != expected]
